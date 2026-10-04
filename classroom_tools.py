"""Доступ до Google Classroom, Drive і Sheets.

Тут немає жодного виклику, що повертає роботу студенту, змінює assignedGrade,
завдання, терміни чи максимальну оцінку. Єдина дозволена зміна в Classroom —
чернетка оцінки (draftGrade); її додамо на етапі 2, після підтвердження викладача.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload

RETRIES = 5
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 60_000

GOOGLE_DOC = "application/vnd.google-apps.document"
GOOGLE_SHEET = "application/vnd.google-apps.spreadsheet"
GOOGLE_SLIDES = "application/vnd.google-apps.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
CODE_EXT = {".txt", ".md", ".py", ".js", ".ts", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".php",
            ".html", ".htm", ".css", ".sql", ".json", ".xml", ".csv", ".kt", ".go", ".rb", ".sh",
            ".ipynb", ".pas", ".vb", ".yaml", ".yml"}


class ApiError(RuntimeError):
    """Зрозуміла викладачу помилка доступу до Google."""


def explain_http_error(e: HttpError, what: str) -> str:
    status = getattr(e.resp, "status", "?")
    body = str(e)
    if "SERVICE_DISABLED" in body or "has not been used in project" in body or "is disabled" in body:
        api = "Google Sheets API" if "sheets" in body.lower() else "потрібний API"
        return (f"{what}: {api} не ввімкнений у проєкті Google Cloud. Відкрийте "
                "https://console.cloud.google.com/apis/library/sheets.googleapis.com "
                "(для Sheets) у проєкті з credentials.json, натисніть «Увімкнути» і зачекайте кілька хвилин.")
    if status == 401:
        return f"{what}: авторизація недійсна. Видаліть token.json і запустіть python auth.py"
    if status == 403:
        return f"{what}: доступ заборонено (403). {_reason(e)}"
    if status == 404:
        return f"{what}: не знайдено (404)."
    if status == 429:
        return f"{what}: перевищено ліміт запитів Google, спробуйте пізніше."
    return f"{what}: помилка Google API {status}. {_reason(e)}"


def _reason(e: HttpError) -> str:
    try:
        return e.error_details[0].get("message", "") if e.error_details else e.reason or ""
    except Exception:  # noqa: BLE001
        return ""


def call(request, what: str):
    """Виконує запит із повторами при тимчасових помилках (429, 5xx, мережа)."""
    try:
        return request.execute(num_retries=RETRIES)
    except HttpError as e:
        raise ApiError(explain_http_error(e, what)) from e


@dataclass
class Services:
    classroom: object
    drive: object
    sheets: object


def build_services(creds) -> Services:
    return Services(
        classroom=build("classroom", "v1", credentials=creds, cache_discovery=False),
        drive=build("drive", "v3", credentials=creds, cache_discovery=False),
        sheets=build("sheets", "v4", credentials=creds, cache_discovery=False),
    )


def _paged(method, key: str, what: str, **kwargs) -> list:
    out, token = [], None
    while True:
        resp = call(method(pageToken=token, **kwargs), what)
        out.extend(resp.get(key, []))
        token = resp.get("nextPageToken")
        if not token:
            return out


# ---------- Classroom: читання ----------

def list_courses(svc: Services) -> list[dict]:
    return _paged(svc.classroom.courses().list, "courses", "Список курсів",
                  teacherId="me", courseStates=["ACTIVE"], pageSize=100)


def _norm(s: str | None) -> str:
    """Порівняння назв: Unicode NFC (у Classroom «й» буває розкладеним на «и» + знак),
    однакові апострофи, пробіли і регістр."""
    s = unicodedata.normalize("NFC", s or "")
    s = re.sub(r"[’ʼ`]", "'", s)
    return re.sub(r"\s+", " ", s).strip().lower()


def find_course(courses: list[dict], name: str, section: str | None = None,
                course_id: str | None = None) -> dict:
    if course_id:
        for c in courses:
            if c["id"] == str(course_id):
                return c
        raise ApiError(f"Курс з id {course_id} не знайдено серед ваших активних курсів")
    matches = [c for c in courses if _norm(c.get("name")) == _norm(name)]
    if section:
        matches = [c for c in matches if _norm(c.get("section")) == _norm(section)]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ApiError(f"Курс «{name}» ({section or 'без групи'}) не знайдено. "
                       "Перевірте назву: python review.py courses")
    raise ApiError(f"Знайдено кілька курсів «{name}». Вкажіть section або course_id у config.yaml")


def list_coursework(svc: Services, course_id: str) -> list[dict]:
    return _paged(svc.classroom.courses().courseWork().list, "courseWork", "Список завдань",
                  courseId=course_id, courseWorkStates=["PUBLISHED"], pageSize=100)


def find_coursework(items: list[dict], title: str) -> dict:
    matches = [c for c in items if _norm(c.get("title")) == _norm(title)]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ApiError(f"Завдання «{title}» не знайдено в курсі. Перелік: python review.py assignments")
    raise ApiError(f"Кілька завдань з назвою «{title}» — перейменуйте одне з них у Classroom")


def list_students(svc: Services, course_id: str) -> dict[str, str]:
    items = _paged(svc.classroom.courses().students().list, "students", "Список студентів",
                   courseId=course_id, pageSize=100)
    return {s["userId"]: s.get("profile", {}).get("name", {}).get("fullName", "") for s in items}


def list_submissions(svc: Services, course_id: str, coursework_id: str) -> list[dict]:
    return _paged(svc.classroom.courses().courseWork().studentSubmissions().list,
                  "studentSubmissions", "Список робіт студентів",
                  courseId=course_id, courseWorkId=coursework_id, pageSize=100)


def turned_in_time(sub: dict) -> str | None:
    """Час останньої здачі з історії роботи (повторна здача дає новий час)."""
    times = [h["stateHistory"]["stateTimestamp"] for h in sub.get("submissionHistory", [])
             if h.get("stateHistory", {}).get("state") == "TURNED_IN"]
    return max(times) if times else sub.get("updateTime")


def is_turned_in(sub: dict) -> bool:
    """Здана (вчасно або із запізненням — late=true), ще не повернута."""
    return sub.get("state") == "TURNED_IN"


def format_time(iso: str | None, tz: str) -> str:
    if not iso:
        return ""
    dt = datetime.fromisoformat(iso.replace("Z", "+00:00"))
    return dt.astimezone(ZoneInfo(tz)).strftime("%d.%m.%Y %H:%M")


# ---------- Drive: вкладення і витяг тексту ----------

@dataclass
class Extracted:
    name: str
    status: str                 # "ok" | "image" | "pdf_scan" | "error"
    text: str = ""
    reason: str = ""            # для status="error": файл недоступний / порожній / непідтримуваний формат
    file_path: str = ""         # для зображень і сканів: локальна копія, яку модель може переглянути
    link: str = ""
    raw: bytes = b""            # для хешу вмісту
    meta: dict = field(default_factory=dict)


def _download(request) -> bytes:
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = dl.next_chunk(num_retries=RETRIES)
        if buf.tell() > MAX_FILE_BYTES:
            raise ValueError("файл завеликий")
    return buf.getvalue()


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Аркуш: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"## Слайд {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return "\n".join(f"## Сторінка {i}\n{p.extract_text() or ''}" for i, p in enumerate(r.pages, 1))


def _safe_name(name: str) -> str:
    return re.sub(r"[^\w.\-]+", "_", name)[:80]


def file_meta(svc: Services, file_id: str) -> dict:
    return call(svc.drive.files().get(
        fileId=file_id, supportsAllDrives=True,
        fields="id,name,mimeType,size,createdTime,modifiedTime,md5Checksum,webViewLink"),
        "Метадані файлу")


def revisions_summary(svc: Services, file_id: str) -> dict:
    """Історія версій (для ознак «вставлено одним фрагментом»). Може бути недоступна — це не помилка."""
    try:
        revs = call(svc.drive.revisions().list(fileId=file_id, fields="revisions(id,modifiedTime)",
                                               pageSize=200), "Історія версій").get("revisions", [])
    except ApiError:
        return {"available": False}
    times = sorted(r["modifiedTime"] for r in revs)
    return {"available": True, "count": len(times),
            "first": times[0] if times else None, "last": times[-1] if times else None}


def extract_drive_file(svc: Services, file_id: str, title: str, save_dir: Path) -> Extracted:
    try:
        meta = file_meta(svc, file_id)
    except ApiError:
        return Extracted(title, "error", reason="файл недоступний")
    name, mime = meta.get("name", title), meta.get("mimeType", "")
    ex = Extracted(name, "ok", link=meta.get("webViewLink", ""),
                   meta={k: meta.get(k) for k in ("mimeType", "createdTime", "modifiedTime")})
    if int(meta.get("size") or 0) > MAX_FILE_BYTES:
        ex.status, ex.reason = "error", "файл завеликий"
        return ex
    ext = Path(name).suffix.lower()
    try:
        if mime == GOOGLE_DOC:
            ex.raw = _download(svc.drive.files().export_media(fileId=file_id, mimeType="text/plain"))
            ex.text = ex.raw.decode("utf-8", "replace")
        elif mime == GOOGLE_SHEET:
            ex.raw = _download(svc.drive.files().export_media(fileId=file_id, mimeType=XLSX))
            ex.text = _xlsx_text(ex.raw)
        elif mime == GOOGLE_SLIDES:
            ex.raw = _download(svc.drive.files().export_media(fileId=file_id, mimeType="text/plain"))
            ex.text = ex.raw.decode("utf-8", "replace")
        elif mime.startswith("application/vnd.google-apps"):
            ex.status, ex.reason = "error", "непідтримуваний формат"
            return ex
        else:
            ex.raw = _download(svc.drive.files().get_media(fileId=file_id, supportsAllDrives=True))
            if mime == "application/pdf" or ext == ".pdf":
                ex.text = _pdf_text(ex.raw)
                if len(re.sub(r"## Сторінка \d+|\s", "", ex.text)) < 50:
                    # Скан без текстового шару: модель може переглянути PDF сама
                    ex.status, ex.text = "pdf_scan", ""
                    ex.file_path = _save(save_dir, file_id, name, ex.raw)
            elif mime == DOCX or ext == ".docx":
                ex.text = _docx_text(ex.raw)
            elif mime == XLSX or ext == ".xlsx":
                ex.text = _xlsx_text(ex.raw)
            elif mime == PPTX or ext == ".pptx":
                ex.text = _pptx_text(ex.raw)
            elif mime.startswith("image/"):
                ex.status = "image"
                ex.file_path = _save(save_dir, file_id, name, ex.raw)
            elif mime.startswith("text/") or ext in CODE_EXT:
                ex.text = ex.raw.decode("utf-8", "replace")
            else:
                ex.status, ex.reason = "error", "непідтримуваний формат"
                return ex
    except ApiError:
        ex.status, ex.reason = "error", "файл недоступний"
        return ex
    except (ValueError, zipfile.BadZipFile, KeyError, OSError) as e:
        ex.status, ex.reason = "error", ("файл завеликий" if "завеликий" in str(e) else "файл пошкоджений або не читається")
        return ex

    if ex.status == "ok":
        if not ex.text.strip():
            ex.status, ex.reason = "error", "порожній"
        elif len(ex.text) > MAX_TEXT_CHARS:
            ex.text = ex.text[:MAX_TEXT_CHARS] + "\n[… текст обрізано через обсяг …]"
    if mime in (GOOGLE_DOC, GOOGLE_SHEET, GOOGLE_SLIDES):
        ex.meta["revisions"] = revisions_summary(svc, file_id)
    return ex


def _save(save_dir: Path, file_id: str, name: str, data: bytes) -> str:
    save_dir.mkdir(parents=True, exist_ok=True)
    p = save_dir / f"{file_id[:12]}_{_safe_name(name)}"
    p.write_bytes(data)
    return str(p)


def extract_submission(svc: Services, sub: dict, save_dir: Path) -> list[Extracted]:
    """Усі вкладення роботи студента + відповідь на питання (якщо завдання-питання)."""
    out: list[Extracted] = []
    short = sub.get("shortAnswerSubmission", {}).get("answer")
    if short:
        out.append(Extracted("Відповідь у Classroom", "ok", text=short, raw=short.encode()))
    mc = sub.get("multipleChoiceSubmission", {}).get("answer")
    if mc:
        out.append(Extracted("Вибрана відповідь", "ok", text=mc, raw=mc.encode()))
    for att in sub.get("assignmentSubmission", {}).get("attachments", []):
        if "driveFile" in att:
            f = att["driveFile"]
            out.append(extract_drive_file(svc, f["id"], f.get("title", "файл"), save_dir))
        elif "link" in att:
            url = att["link"].get("url", "")
            m = re.search(r"/(?:document|spreadsheets|presentation|file)/d/([\w-]{20,})", url)
            if m and "google.com" in url:
                out.append(extract_drive_file(svc, m.group(1), att["link"].get("title", url), save_dir))
            else:
                out.append(Extracted(att["link"].get("title") or url, "error", link=url,
                                     reason="непідтримуваний формат (зовнішнє посилання)"))
        elif "youTubeVideo" in att:
            out.append(Extracted(att["youTubeVideo"].get("title", "відео"), "error",
                                 reason="непідтримуваний формат (відео)"))
        elif "form" in att:
            out.append(Extracted(att["form"].get("title", "форма"), "error",
                                 reason="непідтримуваний формат (Google Форма)"))
    return out


def coursework_task_text(svc: Services, cw: dict, save_dir: Path) -> str:
    """Умова завдання: опис + текст прикріплених матеріалів (Docs/PDF/docx)."""
    parts = [f"Назва: {cw.get('title', '')}", cw.get("description", "")]
    if cw.get("maxPoints"):
        parts.append(f"Максимальна оцінка в Classroom: {cw['maxPoints']}")
    for m in cw.get("materials", []):
        if "driveFile" in m:
            f = m["driveFile"]["driveFile"]
            ex = extract_drive_file(svc, f["id"], f.get("title", "матеріал"), save_dir)
            if ex.status == "ok":
                parts.append(f"### Матеріал завдання: {ex.name}\n{ex.text[:20000]}")
        elif "link" in m:
            parts.append(f"Посилання в завданні: {m['link'].get('title', '')} {m['link'].get('url', '')}")
    return "\n\n".join(p for p in parts if p)


# ---------- Google Sheets: звіт ----------

HEADERS = ["ПІБ студента", "Завдання", "Час здачі", "Статус", "Запропонована оцінка",
           "Бали (зі 100)", "Розбивка за критеріями", "Зауваження", "Ризик ШІ",
           "Примітка щодо ШІ", "Уточнювальні запитання", "Статус чернетки в Classroom",
           "Дата перевірки", "Посилання на роботу", "Рішення викладача", "Ключ (службове)"]
RISK_COL = "I"
KEY_COL_INDEX = len(HEADERS) - 1
SUMMARY_TITLE = "Підсумок"
SUMMARY_HEADERS = ["Група / курс", "Завдань у перевірці", "Здано", "Перевірено агентом",
                   "Не здали", "Середня оцінка (4-бальна)", "Середній бал (зі 100)",
                   "Робіт із середнім/високим ризиком ШІ"]


def sheet_title(course: dict) -> str:
    t = course.get("name", "Курс")
    if course.get("section"):
        t += f" ({course['section']})"
    return re.sub(r"[\[\]:*?/\\]", "-", t)[:95]


def create_spreadsheet(svc: Services, title: str) -> str:
    resp = call(svc.sheets.spreadsheets().create(
        body={"properties": {"title": title, "locale": "uk_UA", "timeZone": "Europe/Kyiv"},
              "sheets": [{"properties": {"title": SUMMARY_TITLE}}]},
        fields="spreadsheetId"), "Створення таблиці звіту")
    ssid = resp["spreadsheetId"]
    call(svc.sheets.spreadsheets().values().update(
        spreadsheetId=ssid, range=f"'{SUMMARY_TITLE}'!A1", valueInputOption="RAW",
        body={"values": [SUMMARY_HEADERS]}), "Заголовки підсумку")
    return ssid


def _sheets_meta(svc: Services, ssid: str) -> dict:
    return call(svc.sheets.spreadsheets().get(
        spreadsheetId=ssid, fields="sheets(properties(sheetId,title))"), "Читання таблиці звіту")


def ensure_course_sheet(svc: Services, ssid: str, title: str) -> int:
    """Створює аркуш групи (заголовки, закріплений рядок, умовне форматування), якщо його ще немає."""
    for s in _sheets_meta(svc, ssid).get("sheets", []):
        if s["properties"]["title"] == title:
            return s["properties"]["sheetId"]
    resp = call(svc.sheets.spreadsheets().batchUpdate(spreadsheetId=ssid, body={"requests": [
        {"addSheet": {"properties": {"title": title, "gridProperties": {"frozenRowCount": 1}}}}]}),
        "Створення аркуша групи")
    sid = resp["replies"][0]["addSheet"]["properties"]["sheetId"]
    call(svc.sheets.spreadsheets().values().update(
        spreadsheetId=ssid, range=f"'{title}'!A1", valueInputOption="RAW",
        body={"values": [HEADERS]}), "Заголовки аркуша")
    rng = {"sheetId": sid, "startRowIndex": 1, "startColumnIndex": 0, "endColumnIndex": len(HEADERS)}

    def rule(value, rgb, index):
        return {"addConditionalFormatRule": {"index": index, "rule": {"ranges": [rng], "booleanRule": {
            "condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": f'=${RISK_COL}2="{value}"'}]},
            "format": {"backgroundColor": rgb}}}}}

    call(svc.sheets.spreadsheets().batchUpdate(spreadsheetId=ssid, body={"requests": [
        rule("високий", {"red": 0.96, "green": 0.78, "blue": 0.76}, 0),
        rule("середній", {"red": 1.0, "green": 0.93, "blue": 0.7}, 1),
        {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 0, "endRowIndex": 1},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
                        "fields": "userEnteredFormat.textFormat.bold"}},
        {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 1, "startColumnIndex": 6, "endColumnIndex": 11},
                        "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
                        "fields": "userEnteredFormat(wrapStrategy,verticalAlignment)"}},
        {"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS",
                                                 "startIndex": 6, "endIndex": 11},
                                       "properties": {"pixelSize": 320}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS",
                                                 "startIndex": KEY_COL_INDEX, "endIndex": KEY_COL_INDEX + 1},
                                       "properties": {"hiddenByUser": True}, "fields": "hiddenByUser"}},
    ]}), "Форматування аркуша")
    return sid


def existing_keys(svc: Services, ssid: str, title: str) -> set[str]:
    col = chr(ord("A") + KEY_COL_INDEX)
    resp = call(svc.sheets.spreadsheets().values().get(
        spreadsheetId=ssid, range=f"'{title}'!{col}2:{col}"), "Читання аркуша")
    return {r[0] for r in resp.get("values", []) if r}


def append_rows(svc: Services, ssid: str, title: str, rows: list[list]) -> None:
    """Лише додає рядки в кінець — наявні рядки (і правки викладача) не змінюються."""
    if rows:
        call(svc.sheets.spreadsheets().values().append(
            spreadsheetId=ssid, range=f"'{title}'!A1", valueInputOption="RAW",
            insertDataOption="INSERT_ROWS", body={"values": rows}), "Додавання рядків у звіт")


def write_summary(svc: Services, ssid: str, rows: list[list]) -> None:
    """Аркуш «Підсумок» обчислюваний, тому перезаписується повністю."""
    titles = {s["properties"]["title"] for s in _sheets_meta(svc, ssid).get("sheets", [])}
    if SUMMARY_TITLE not in titles:
        call(svc.sheets.spreadsheets().batchUpdate(spreadsheetId=ssid, body={"requests": [
            {"addSheet": {"properties": {"title": SUMMARY_TITLE, "index": 0}}}]}), "Аркуш підсумку")
    call(svc.sheets.spreadsheets().values().clear(
        spreadsheetId=ssid, range=f"'{SUMMARY_TITLE}'!A:Z", body={}), "Очищення підсумку")
    call(svc.sheets.spreadsheets().values().update(
        spreadsheetId=ssid, range=f"'{SUMMARY_TITLE}'!A1", valueInputOption="RAW",
        body={"values": [SUMMARY_HEADERS] + rows}), "Запис підсумку")
